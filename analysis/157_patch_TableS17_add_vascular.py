#!/usr/bin/env python3

from pathlib import Path
import hashlib
import shutil
import pandas as pd
from openpyxl import load_workbook

ROOT = Path("/data/qiallab/Framingham")
REPO = ROOT / "FINAL_MANUSCRIPT_REPO_20260928"

OUTDIR = REPO / "SupplementaryTables" / "TableS17"
CODEDIR = OUTDIR / "TableS17_code"
SRCDIR = OUTDIR / "TableS17_source_data"

BASE_XLSX = (
    OUTDIR /
    "TableS17_LongitudinalRobustness_BEFORE_VASCULAR_ROWS.xlsx"
)

FINAL_XLSX = OUTDIR / "TableS17_LongitudinalRobustness.xlsx"
FINAL_CSV = OUTDIR / "TableS17_LongitudinalRobustness.csv"
README = OUTDIR / "README_PROVENANCE.txt"

VASC_SRC = (
    ROOT / "results" /
    "longitudinal_heart_brain" /
    "temporal_prediction" /
    "vascular_adjustment" /
    "LA_ventricle_early_vascular_adjustment.tsv"
)

VASC_SCRIPT = (
    ROOT / "code" /
    "111_longitudinal_LA_ventricle_vascular_adjustment.py"
)

PATCH_SCRIPT = (
    ROOT / "code" /
    "157_patch_TableS17_add_vascular.py"
)

print("=" * 110)
print("PATCH TABLE S17 WITH VERIFIED EARLY-VASCULAR SENSITIVITIES")
print("=" * 110)

for p in [BASE_XLSX, VASC_SRC, VASC_SCRIPT]:
    if not p.exists():
        raise FileNotFoundError(p)

CODEDIR.mkdir(parents=True, exist_ok=True)
SRCDIR.mkdir(parents=True, exist_ok=True)

# ------------------------------------------------------------------
# Preserve malformed/current files before replacing them
# ------------------------------------------------------------------

for src, name in [
    (
        FINAL_XLSX,
        "TableS17_LongitudinalRobustness_BEFORE_PROVENANCE_REPAIR.xlsx",
    ),
    (
        FINAL_CSV,
        "TableS17_LongitudinalRobustness_BEFORE_PROVENANCE_REPAIR.csv",
    ),
]:
    dst = OUTDIR / name
    if src.exists() and not dst.exists():
        shutil.copy2(src, dst)

# ------------------------------------------------------------------
# Read authoritative vascular results
# ------------------------------------------------------------------

v = pd.read_csv(VASC_SRC, sep="\t")

want = v[
    v["subset"].eq("prospective")
    & v["term"].eq("MRI_time_x_LA_x_APOE4")
    & v["model"].isin([
        "M1_early_vascular",
        "M2_early_vascular_plus_DBP",
    ])
].copy()

if len(want) != 2:
    raise RuntimeError(
        f"Expected exactly 2 vascular interaction rows; found {len(want)}"
    )

order = [
    "M1_early_vascular",
    "M2_early_vascular_plus_DBP",
]

want["__order"] = want["model"].map(
    {m: i for i, m in enumerate(order)}
)
want = want.sort_values("__order")

labels = {
    "M1_early_vascular":
        "+ early vascular covariates",
    "M2_early_vascular_plus_DBP":
        "+ early vascular covariates + DBP",
}

print("\nAUTHORITATIVE VASCULAR ROWS:")
print(
    want[
        [
            "subset", "model", "term",
            "beta", "SE", "CI_low", "CI_high",
            "P", "N_subjects", "N_observations",
            "converged",
        ]
    ].to_string(index=False)
)

# ------------------------------------------------------------------
# Start from intact 3-sheet pre-vascular workbook
# ------------------------------------------------------------------

wb = load_workbook(BASE_XLSX)

expected_sheets = ["Results", "Numeric_Source", "Provenance"]

if wb.sheetnames != expected_sheets:
    raise RuntimeError(
        f"Unexpected base workbook sheets: {wb.sheetnames}"
    )

wsr = wb["Results"]
wsn = wb["Numeric_Source"]
wsp = wb["Provenance"]

# Remove vascular rows if script is rerun
vascular_labels = set(labels.values())

for ws in [wsr, wsn]:
    delete_rows = []
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=1).value in vascular_labels:
            delete_rows.append(row)

    for row in reversed(delete_rows):
        ws.delete_rows(row)

# ------------------------------------------------------------------
# Append Results rows
# ------------------------------------------------------------------

for _, r in want.iterrows():
    wsr.append([
        labels[r["model"]],
        int(r["N_subjects"]),
        int(r["N_observations"]),
        f'{float(r["beta"]):.5f}',
        f'{float(r["CI_low"]):.5f} to {float(r["CI_high"]):.5f}',
        f'{float(r["P"]):.2e}' if float(r["P"]) < 0.001
        else f'{float(r["P"]):.4f}',
    ])

# ------------------------------------------------------------------
# Append Numeric_Source rows at full precision
# ------------------------------------------------------------------

for _, r in want.iterrows():
    wsn.append([
        labels[r["model"]],
        int(r["N_subjects"]),
        int(r["N_observations"]),
        float(r["beta"]),
        float(r["CI_low"]),
        float(r["CI_high"]),
        float(r["P"]),
        (
            "111_longitudinal_LA_ventricle_"
            "vascular_adjustment.py"
            f" [{r['model']}]"
        ),
    ])

# ------------------------------------------------------------------
# Add provenance entries
# ------------------------------------------------------------------

existing_prov = {
    str(wsp.cell(row=i, column=1).value)
    for i in range(2, wsp.max_row + 1)
}

prov_items = [
    (
        "early_vascular_adjustment_script",
        VASC_SCRIPT,
    ),
    (
        "early_vascular_adjustment_effect_file",
        VASC_SRC,
    ),
    (
        "TableS17_vascular_patch_script",
        PATCH_SCRIPT,
    ),
]

for analysis, path in prov_items:
    if analysis in existing_prov:
        continue

    sha = hashlib.sha256(path.read_bytes()).hexdigest()

    wsp.append([
        analysis,
        str(path),
        sha,
    ])

# ------------------------------------------------------------------
# Save workbook
# ------------------------------------------------------------------

wb.save(FINAL_XLSX)

# ------------------------------------------------------------------
# Rebuild publication CSV directly from Results sheet
# ------------------------------------------------------------------

wb_check = load_workbook(FINAL_XLSX, data_only=False)
ws = wb_check["Results"]

rows = list(ws.values)
header = list(rows[0])
data = rows[1:]

pub = pd.DataFrame(data, columns=header)
pub.to_csv(FINAL_CSV, index=False)

# ------------------------------------------------------------------
# Freeze source/script artifacts
# ------------------------------------------------------------------

shutil.copy2(
    VASC_SCRIPT,
    CODEDIR / VASC_SCRIPT.name
)

shutil.copy2(
    VASC_SRC,
    SRCDIR / VASC_SRC.name
)

shutil.copy2(
    PATCH_SCRIPT,
    CODEDIR / PATCH_SCRIPT.name
)

# ------------------------------------------------------------------
# Update README
# ------------------------------------------------------------------

readme_text = README.read_text() if README.exists() else ""

vascular_section = """
EARLY VASCULAR COVARIATE SENSITIVITY
------------------------------------
111: longitudinal LA-ventricle early vascular adjustment.

Exam-4 vascular covariates were used rather than later Exam-8
covariates. The archived sensitivity models include:

M1_early_vascular:
    prospective beta = -0.008821
    95% CI = -0.013329 to -0.004313
    P = 0.000125
    N = 1,269 subjects
    MRI observations = 3,685

M2_early_vascular_plus_DBP:
    prospective beta = -0.008818
    95% CI = -0.013326 to -0.004310
    P = 0.000126
    N = 1,269 subjects
    MRI observations = 3,685

These replace the previously unverified longitudinal statement
describing adjustment for CMR-derived LV mass index and LVESVi.

Canonical source:
results/longitudinal_heart_brain/temporal_prediction/
vascular_adjustment/LA_ventricle_early_vascular_adjustment.tsv

"""

if "EARLY VASCULAR COVARIATE SENSITIVITY" not in readme_text:
    marker = "EXCLUDED ANALYSES\n-----------------\n"

    if marker in readme_text:
        readme_text = readme_text.replace(
            marker,
            vascular_section + marker
        )
    else:
        readme_text += "\n" + vascular_section

README.write_text(readme_text)

# ------------------------------------------------------------------
# Validation
# ------------------------------------------------------------------

wbv = load_workbook(FINAL_XLSX, data_only=False)

assert wbv.sheetnames == [
    "Results",
    "Numeric_Source",
    "Provenance",
]

res = pd.read_excel(
    FINAL_XLSX,
    sheet_name="Results",
)

num = pd.read_excel(
    FINAL_XLSX,
    sheet_name="Numeric_Source",
)

prov = pd.read_excel(
    FINAL_XLSX,
    sheet_name="Provenance",
)

assert len(res) == 10
assert len(num) == 10

assert res.iloc[-2]["Model / sensitivity analysis"] == (
    "+ early vascular covariates"
)
assert res.iloc[-1]["Model / sensitivity analysis"] == (
    "+ early vascular covariates + DBP"
)

assert int(res.iloc[-2]["N"]) == 1269
assert int(res.iloc[-2]["MRI observations"]) == 3685

assert abs(
    float(num.iloc[-2]["Beta interaction"])
    - (-0.008821224779947654)
) < 1e-12

assert abs(
    float(num.iloc[-1]["Beta interaction"])
    - (-0.008818389142354058)
) < 1e-12

print("\n" + "=" * 110)
print("FINAL RESULTS SHEET")
print("=" * 110)
print(res.to_string(index=False))

print("\n" + "=" * 110)
print("PROVENANCE — LAST ROWS")
print("=" * 110)
print(prov.tail(5).to_string(index=False))

print("\nWorkbook sheets:", wbv.sheetnames)

print("\nFILES:")
print(FINAL_CSV)
print(FINAL_XLSX)
print(README)

print("\nFROZEN VASCULAR FILES:")
print(CODEDIR / VASC_SCRIPT.name)
print(SRCDIR / VASC_SRC.name)
print(CODEDIR / PATCH_SCRIPT.name)

print(
    "\nSUCCESS: Table S17 is numerically correct, "
    "three-sheet structure restored, and vascular provenance frozen."
)
