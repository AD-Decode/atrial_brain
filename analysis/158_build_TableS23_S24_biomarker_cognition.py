#!/usr/bin/env python3

from pathlib import Path
import shutil
import hashlib
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ROOT = Path("/data/qiallab/Framingham")
REPO = ROOT / "FINAL_MANUSCRIPT_REPO_20260928"
SUPP = REPO / "SupplementaryTables"

S23 = SUPP / "TableS23"
S24 = SUPP / "TableS24"

for d in [S23, S24]:
    d.mkdir(parents=True, exist_ok=True)

# ============================================================
# SOURCE FILES
# ============================================================

BIO = (
    ROOT / "results/longitudinal_heart_brain/"
    "plasma_biomarker_prediction/multicardiac_biomarkers/"
    "cardiac_trajectory_APOE4_interactions_biomarkers.tsv"
)

COG = (
    ROOT / "results/longitudinal_heart_brain/"
    "cognition_prediction/"
    "cardiac_trajectory_later_cognition_APOE4_interactions.tsv"
)

MMSE = (
    ROOT / "results/longitudinal_heart_brain/"
    "cognition_prediction/longitudinal_MMSE/"
    "MMSE_cardiac_trajectory_APOE4_slope_interactions.tsv"
)

COG_DOMAIN = (
    ROOT / "results/longitudinal_heart_brain/"
    "cognition_prediction/cognitive_domains/"
    "cardiac_APOE4_later_cognitive_domain_interactions.tsv"
)

SCRIPT = ROOT / "code/158_build_TableS23_S24_biomarker_cognition.py"

for p in [BIO, COG, MMSE]:
    if not p.exists():
        raise FileNotFoundError(p)

# ============================================================
# HELPERS
# ============================================================

def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def autosize(ws):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in col
        )
        ws.column_dimensions[letter].width = min(width + 2, 45)

def style_header(ws):
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"

def write_df(ws, df):
    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append(list(row))
    style_header(ws)
    autosize(ws)

# ============================================================
# TABLE S23 — PLASMA BIOMARKERS
# ============================================================

bio = pd.read_csv(BIO, sep="\t")

expected_bio = [
    "cardiac",
    "cardiac_short",
    "biomarker",
    "N",
    "N_APOE4_noncarrier",
    "N_APOE4_carrier",
    "interaction_beta",
    "interaction_SE",
    "interaction_CI_low",
    "interaction_CI_high",
    "interaction_P",
    "interaction_q_within_cardiac4",
    "interaction_q_global16",
]

if list(bio.columns) != expected_bio:
    raise RuntimeError(
        f"Unexpected biomarker schema:\n{list(bio.columns)}"
    )

if len(bio) != 16:
    raise RuntimeError(
        f"Expected 16 biomarker interaction tests; found {len(bio)}"
    )

bio_pub = pd.DataFrame({
    "Cardiac trajectory": bio["cardiac"],
    "Plasma biomarker": bio["biomarker"],
    "N": bio["N"],
    "APOE ε4 noncarriers": bio["N_APOE4_noncarrier"],
    "APOE ε4 carriers": bio["N_APOE4_carrier"],
    "β interaction": bio["interaction_beta"],
    "SE": bio["interaction_SE"],
    "95% CI": [
        f"{lo:.4f} to {hi:.4f}"
        for lo, hi in zip(
            bio["interaction_CI_low"],
            bio["interaction_CI_high"],
        )
    ],
    "P": bio["interaction_P"],
    "q within cardiac family (4 tests)":
        bio["interaction_q_within_cardiac4"],
    "q global (16 tests)":
        bio["interaction_q_global16"],
})

s23_csv = S23 / "TableS23_CardiacTrajectories_PlasmaBiomarkers.csv"
s23_xlsx = S23 / "TableS23_CardiacTrajectories_PlasmaBiomarkers.xlsx"

bio_pub.to_csv(s23_csv, index=False)

wb = Workbook()
ws = wb.active
ws.title = "Results"
write_df(ws, bio_pub)

ws2 = wb.create_sheet("Numeric_Source")
write_df(ws2, bio)

prov23 = pd.DataFrame([
    {
        "analysis": "plasma biomarker interaction family",
        "source": str(BIO),
        "sha256": sha256(BIO),
    },
    {
        "analysis": "table build script",
        "source": str(SCRIPT),
        "sha256": sha256(SCRIPT),
    },
])

ws3 = wb.create_sheet("Provenance")
write_df(ws3, prov23)

wb.save(s23_xlsx)

# freeze exact source
src23 = S23 / "TableS23_source_data"
code23 = S23 / "TableS23_code"
src23.mkdir(exist_ok=True)
code23.mkdir(exist_ok=True)

shutil.copy2(BIO, src23 / BIO.name)
shutil.copy2(SCRIPT, code23 / SCRIPT.name)

(S23 / "README_PROVENANCE.txt").write_text(
"""Table S23. Associations of antecedent cardiac trajectories with later plasma biomarkers by APOE ε4 status.

The table contains the complete 4 cardiac trajectory × 4 plasma biomarker interaction family (16 tests).

Cardiac trajectories:
- LA dimension
- LV mass
- fractional shortening
- LV end-diastolic dimension

Biomarkers:
- Aβ42/40
- p-tau181
- GFAP
- NfL

Both within-cardiac four-test FDR and global 16-test FDR values are reported.

No interaction survives FDR correction in the global 16-test family.

The workbook contains Results, Numeric_Source, and Provenance sheets.
"""
)

# ============================================================
# TABLE S24 — COGNITION
# ============================================================

cog = pd.read_csv(COG, sep="\t")
mmse = pd.read_csv(MMSE, sep="\t")

# publication-facing prospective subset only
cog_pro = cog[cog["subset"].eq("prospective")].copy()
mmse_pro = mmse[mmse["subset"].eq("prospective")].copy()

if len(cog_pro) != 16:
    raise RuntimeError(
        f"Expected 16 prospective later-cognition tests; found {len(cog_pro)}"
    )

if len(mmse_pro) != 4:
    raise RuntimeError(
        f"Expected 4 prospective MMSE tests; found {len(mmse_pro)}"
    )

cog_pub = pd.DataFrame({
    "Cardiac trajectory": cog_pro["cardiac"],
    "Cognitive outcome": cog_pro["cognition"],
    "N": cog_pro["N"],
    "β interaction": cog_pro["interaction_beta"],
    "SE": cog_pro["interaction_SE"],
    "95% CI": [
        f"{lo:.4f} to {hi:.4f}"
        for lo, hi in zip(
            cog_pro["interaction_CI_low"],
            cog_pro["interaction_CI_high"],
        )
    ],
    "P": cog_pro["interaction_P"],
    "q global (16 tests)": cog_pro["interaction_q_global16"],
    "q within cardiac family (4 tests)":
        cog_pro["interaction_q_within_cardiac4"],
})

mmse_pub = pd.DataFrame({
    "Cardiac trajectory": mmse_pro["cardiac"],
    "N subjects": mmse_pro["N_subjects"],
    "MMSE observations": mmse_pro["N_observations"],
    "β cardiac trajectory × cognitive time × APOE ε4":
        mmse_pro["time_x_cardiac_x_APOE4_beta"],
    "SE": mmse_pro["SE"],
    "95% CI": [
        f"{lo:.4f} to {hi:.4f}"
        for lo, hi in zip(
            mmse_pro["CI_low"],
            mmse_pro["CI_high"],
        )
    ],
    "P": mmse_pro["P"],
    "q FDR (4 tests)": mmse_pro["q_FDR4"],
})

s24a = S24 / "TableS24A_CardiacTrajectories_LaterCognition.csv"
s24b = S24 / "TableS24B_CardiacTrajectories_LongitudinalMMSE.csv"
s24x = S24 / "TableS24_Cognition.xlsx"

cog_pub.to_csv(s24a, index=False)
mmse_pub.to_csv(s24b, index=False)

wb = Workbook()

ws = wb.active
ws.title = "S24A_LaterCognition"
write_df(ws, cog_pub)

ws2 = wb.create_sheet("S24B_LongitudinalMMSE")
write_df(ws2, mmse_pub)

# exact full sources, including gap ≥2y duplicates
ws3 = wb.create_sheet("Source_LaterCognition")
write_df(ws3, cog)

ws4 = wb.create_sheet("Source_MMSE")
write_df(ws4, mmse)

prov_rows = [
    {
        "analysis": "later cognition APOE4 interactions",
        "source": str(COG),
        "sha256": sha256(COG),
    },
    {
        "analysis": "longitudinal MMSE APOE4 interactions",
        "source": str(MMSE),
        "sha256": sha256(MMSE),
    },
    {
        "analysis": "table build script",
        "source": str(SCRIPT),
        "sha256": sha256(SCRIPT),
    },
]

# include cognitive-domain source in provenance if available,
# but do not mix exploratory domains into publication-facing table
if COG_DOMAIN.exists():
    prov_rows.append({
        "analysis": "additional cognitive-domain interaction output",
        "source": str(COG_DOMAIN),
        "sha256": sha256(COG_DOMAIN),
    })

ws5 = wb.create_sheet("Provenance")
write_df(ws5, pd.DataFrame(prov_rows))

wb.save(s24x)

src24 = S24 / "TableS24_source_data"
code24 = S24 / "TableS24_code"
src24.mkdir(exist_ok=True)
code24.mkdir(exist_ok=True)

for f in [COG, MMSE]:
    shutil.copy2(f, src24 / f.name)

if COG_DOMAIN.exists():
    shutil.copy2(COG_DOMAIN, src24 / COG_DOMAIN.name)

shutil.copy2(SCRIPT, code24 / SCRIPT.name)

(S24 / "README_PROVENANCE.txt").write_text(
"""Table S24. Associations of antecedent cardiac trajectories with later cognition by APOE ε4 status.

Panel A:
Prospective analyses of four cardiac trajectories with four later cognitive outcomes (16 tests):
- Logical Memory delayed
- Trails B
- Animal fluency
- Digit Span Backward

Both global 16-test FDR and within-cardiac four-test FDR are shown.

The LV end-diastolic dimension × APOE ε4 interaction for Digit Span Backward has nominal P=0.0112 and within-cardiac q=0.0448, but global q=0.179; it therefore does not survive correction across the complete 16-test cognition family.

Panel B:
Prospective longitudinal MMSE mixed-effects interaction analyses for the four cardiac trajectories.

LV mass trajectory × cognitive time × APOE ε4:
beta=-0.007859
P=0.00731
q=0.02925 across the four MMSE tests.

The >=2-year-gap rows are retained in the source sheets for auditability but are not duplicated in the publication-facing panels because the archived values are identical to the prospective rows in these cognition outputs.

The workbook contains publication-facing panels, full source sheets, and provenance.
"""
)

# ============================================================
# VALIDATION
# ============================================================

print("=" * 115)
print("TABLE S23")
print("=" * 115)
print(bio_pub.to_string(index=False))

print("\nMinimum global biomarker q:",
      bio["interaction_q_global16"].min())

print("\n" + "=" * 115)
print("TABLE S24A — LATER COGNITION")
print("=" * 115)
print(cog_pub.to_string(index=False))

print("\nMinimum global cognition q:",
      cog_pro["interaction_q_global16"].min())

print("\n" + "=" * 115)
print("TABLE S24B — LONGITUDINAL MMSE")
print("=" * 115)
print(mmse_pub.to_string(index=False))

lv = mmse_pro[mmse_pro["cardiac"].eq("LV mass")]

assert len(lv) == 1
assert abs(float(lv.iloc[0]["P"]) - 0.007312) < 1e-5
assert abs(float(lv.iloc[0]["q_FDR4"]) - 0.029247) < 1e-5

assert len(bio_pub) == 16
assert len(cog_pub) == 16
assert len(mmse_pub) == 4

print("\nFILES CREATED:")
for p in [
    s23_csv, s23_xlsx,
    s24a, s24b, s24x,
]:
    print(p)

print(
    "\nSUCCESS: Tables S23 and S24 created with "
    "publication-facing results, full numeric sources, and provenance."
)
