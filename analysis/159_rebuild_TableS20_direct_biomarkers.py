#!/usr/bin/env python3

from pathlib import Path
import hashlib
import shutil
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ROOT = Path("/data/qiallab/Framingham")
REPO = ROOT / "FINAL_MANUSCRIPT_REPO_20260928"
OUT = REPO / "SupplementaryTables" / "TableS20"

OUT.mkdir(parents=True, exist_ok=True)

CODEDIR = OUT / "TableS20_code"
SRCDIR = OUT / "TableS20_source_data"

CODEDIR.mkdir(exist_ok=True)
SRCDIR.mkdir(exist_ok=True)

RESULT = (
    ROOT / "results/direct_biomarker_prediction/"
    "cardiac_APOE4_biomarker_interactions.csv"
)

ORIG_SCRIPT = (
    ROOT / "code/44_direct_cardiac_APOE4_biomarker_prediction.py"
)

AVAIL = (
    ROOT / "results/direct_biomarker_prediction/"
    "biomarker_prediction_availability.csv"
)

DATASET = (
    ROOT / "results/direct_biomarker_prediction/"
    "cardiac_biomarker_analysis_dataset.csv"
)

MANIFEST = (
    ROOT / "results/direct_biomarker_prediction/"
    "analysis_manifest.json"
)

BUILDER = (
    ROOT / "code/159_rebuild_TableS20_direct_biomarkers.py"
)

print("=" * 115)
print("REBUILD TABLE S20 — DIRECT CARDIAC × APOE4 PLASMA BIOMARKER ANALYSES")
print("=" * 115)

if not RESULT.exists():
    raise FileNotFoundError(RESULT)

df = pd.read_csv(RESULT)

expected = [
    "biomarker",
    "biomarker_column",
    "cardiac",
    "cardiac_column",
    "N",
    "N_APOE4_noncarrier",
    "N_APOE4_carrier",
    "status",
    "interaction_beta",
    "interaction_SE_HC3",
    "interaction_CI_low",
    "interaction_CI_high",
    "interaction_P",
    "slope_noncarrier",
    "slope_noncarrier_SE",
    "slope_noncarrier_P",
    "slope_carrier",
    "slope_carrier_SE",
    "slope_carrier_P",
    "R2",
    "adjusted_R2",
    "interaction_q_BH_16",
    "FDR_significant",
]

missing = [c for c in expected if c not in df.columns]

if missing:
    raise RuntimeError(
        "Missing expected columns:\n" + "\n".join(missing)
    )

if len(df) != 16:
    raise RuntimeError(
        f"Expected complete 16-test family; found {len(df)} rows."
    )

# ------------------------------------------------------------
# Publication-facing table
# ------------------------------------------------------------

pub = pd.DataFrame({
    "Cardiac phenotype": df["cardiac"],
    "Plasma biomarker": df["biomarker"],
    "N": df["N"],
    "APOE ε4 noncarriers": df["N_APOE4_noncarrier"],
    "APOE ε4 carriers": df["N_APOE4_carrier"],
    "β interaction": df["interaction_beta"],
    "HC3 SE": df["interaction_SE_HC3"],
    "95% CI": [
        f"{lo:.4f} to {hi:.4f}"
        for lo, hi in zip(
            df["interaction_CI_low"],
            df["interaction_CI_high"],
        )
    ],
    "P": df["interaction_P"],
    "q BH-FDR (16 tests)": df["interaction_q_BH_16"],
    "FDR significant": df["FDR_significant"],
})

CSV = OUT / "TableS20_DirectCardiac_APOE4_PlasmaBiomarkers.csv"
XLSX = OUT / "TableS20_DirectCardiac_APOE4_PlasmaBiomarkers.xlsx"

pub.to_csv(CSV, index=False)

# ------------------------------------------------------------
# Workbook helpers
# ------------------------------------------------------------

def style(ws):
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"

    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max(
            len(str(x.value)) if x.value is not None else 0
            for x in col
        )
        ws.column_dimensions[letter].width = min(width + 2, 45)

def write_df(ws, x):
    ws.append(list(x.columns))
    for row in x.itertuples(index=False, name=None):
        ws.append(list(row))
    style(ws)

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

# ------------------------------------------------------------
# Build 3-sheet workbook
# ------------------------------------------------------------

wb = Workbook()

ws = wb.active
ws.title = "Results"
write_df(ws, pub)

ws2 = wb.create_sheet("Numeric_Source")
write_df(ws2, df)

prov_rows = []

for analysis, path in [
    ("direct biomarker interaction results", RESULT),
    ("original analysis script", ORIG_SCRIPT),
    ("table build script", BUILDER),
    ("biomarker availability", AVAIL),
    ("analysis dataset", DATASET),
    ("analysis manifest", MANIFEST),
]:
    if path.exists():
        prov_rows.append({
            "analysis": analysis,
            "source": str(path),
            "sha256": sha(path),
        })

prov = pd.DataFrame(prov_rows)

ws3 = wb.create_sheet("Provenance")
write_df(ws3, prov)

wb.save(XLSX)

# ------------------------------------------------------------
# Freeze source files/code
# ------------------------------------------------------------

for p in [RESULT, AVAIL, DATASET, MANIFEST]:
    if p.exists():
        shutil.copy2(p, SRCDIR / p.name)

for p in [ORIG_SCRIPT, BUILDER]:
    if p.exists():
        shutil.copy2(p, CODEDIR / p.name)

# ------------------------------------------------------------
# README
# ------------------------------------------------------------

README = OUT / "README_PROVENANCE.txt"

README.write_text(
"""Supplementary Table S20. Direct cardiac phenotype × APOE ε4 associations with circulating plasma biomarkers.

This table reports the complete prespecified 16-test family from the
cross-sectional biomarker analysis.

Four cardiac phenotypes were tested against four circulating biomarkers:
Aβ42/40, p-tau181, GFAP, and NfL.

Models were adjusted for the prespecified clinical covariates used in the
direct biomarker-prediction analysis and used HC3 robust covariance estimates.

Benjamini-Hochberg FDR correction was applied across the complete 16-test family.

No cardiac × APOE ε4 biomarker interaction survived FDR correction.

This analysis is distinct from:
1. the biomarker-adjustment sensitivity of the four primary cardiac-brain MRI
   interactions; and
2. Table S23, which examines antecedent longitudinal cardiac trajectories in
   relation to later plasma biomarkers.

Workbook sheets:
- Results
- Numeric_Source
- Provenance

The exact analysis outputs and scripts are frozen in the accompanying source
and code directories.
"""
)

# ------------------------------------------------------------
# Validate
# ------------------------------------------------------------

assert len(pub) == 16
assert df["interaction_q_BH_16"].notna().all()
assert not df["FDR_significant"].astype(bool).any()

print("\nPUBLICATION TABLE:")
print(pub.to_string(index=False))

print("\nMinimum nominal P:")
print(df.loc[df["interaction_P"].idxmin(), [
    "cardiac",
    "biomarker",
    "interaction_beta",
    "interaction_P",
    "interaction_q_BH_16",
]].to_string())

print("\nFDR-significant rows:",
      int(df["FDR_significant"].astype(bool).sum()))

print("\nFILES:")
print(CSV)
print(XLSX)
print(README)

print("\nFROZEN SOURCE FILES:")
for p in sorted(SRCDIR.iterdir()):
    print(p)

print("\nFROZEN CODE:")
for p in sorted(CODEDIR.iterdir()):
    print(p)

print(
    "\nSUCCESS: Table S20 rebuilt as the complete "
    "cross-sectional 16-test cardiac × APOE4 biomarker family."
)
