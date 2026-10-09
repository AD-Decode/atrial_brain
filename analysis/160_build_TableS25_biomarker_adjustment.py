#!/usr/bin/env python3

from pathlib import Path
import hashlib
import shutil
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ROOT = Path("/data/qiallab/Framingham")
REPO = ROOT / "FINAL_MANUSCRIPT_REPO_20260928"
OUT = REPO / "SupplementaryTables" / "TableS25"

OUT.mkdir(parents=True, exist_ok=True)

SRCDIR = OUT / "TableS25_source_data"
CODEDIR = OUT / "TableS25_code"

SRCDIR.mkdir(exist_ok=True)
CODEDIR.mkdir(exist_ok=True)

FULL = (
    ROOT / "results/tier3_AD_biomarkers_completecase/"
    "figure_biomarker_robustness_FINAL_table.csv"
)

NUMERIC = (
    ROOT / "results/tier3_AD_biomarkers_completecase/"
    "APOE_AD_biomarker_completecase_proper_CI.csv"
)

SUMMARY = (
    ROOT / "results/tier3_AD_biomarkers_completecase/"
    "APOE_AD_biomarker_completecase_summary.csv"
)

ANALYSIS_SCRIPT = (
    ROOT / "code/40_APOE_AD_completecase_proper_CI.py"
)

FIG_SCRIPT = (
    ROOT / "code/43_plot_FINAL_biomarker_robustness.py"
)

BUILDER = (
    ROOT / "code/160_build_TableS25_biomarker_adjustment.py"
)

for p in [FULL, NUMERIC, SUMMARY, ANALYSIS_SCRIPT]:
    if not p.exists():
        raise FileNotFoundError(p)

print("=" * 120)
print("BUILD TABLE S25 — BIOMARKER ADJUSTMENT OF FOUR PRIMARY CARDIAC–BRAIN INTERACTIONS")
print("=" * 120)

df = pd.read_csv(FULL)
num = pd.read_csv(NUMERIC)
summ = pd.read_csv(SUMMARY)

if len(df) != 28:
    raise RuntimeError(
        f"Expected 28 rows = 4 primary associations × 7 biomarker specs; found {len(df)}"
    )

if len(summ) != 4:
    raise RuntimeError(
        f"Expected 4 summary rows; found {len(summ)}"
    )

# ------------------------------------------------------------
# Human-readable labels
# ------------------------------------------------------------

spec_labels = {
    "M1_AMYLOID": "+ Aβ42/40",
    "M2_PTAU": "+ p-tau181",
    "M3_GFAP": "+ GFAP",
    "M4_NFL": "+ NfL",
    "M5_AMYLOID_PTAU": "+ Aβ42/40 + p-tau181",
    "M6_GFAP_NFL": "+ GFAP + NfL",
    "M7_ALL4": "+ Aβ42/40 + p-tau181 + GFAP + NfL",
}

assoc_labels = {
    ("LVESVi", "rh_superiorfrontal_area"):
        "LVESVi × APOE ε4 → right superior frontal surface area",

    ("LV_MASSi", "rh_lingual_vol"):
        "LV mass index × APOE ε4 → right lingual volume",

    ("LV_MASSi", "rh_lingual_area"):
        "LV mass index × APOE ε4 → right lingual surface area",

    ("LVEF", "lh_middletemporal_vol"):
        "LVEF × APOE ε4 → left middle temporal volume",
}

def assoc_label(row):
    key = (row["cardiac"], row["region"])
    if key not in assoc_labels:
        raise RuntimeError(f"Unexpected association: {key}")
    return assoc_labels[key]

df["Association"] = df.apply(assoc_label, axis=1)
df["Biomarker adjustment"] = df["spec"].map(spec_labels)

if df["Biomarker adjustment"].isna().any():
    raise RuntimeError("Unmapped biomarker specification detected.")

# ------------------------------------------------------------
# Publication-facing table
# ------------------------------------------------------------

pub = pd.DataFrame({
    "Primary cardiac–brain interaction": df["Association"],
    "Biomarker adjustment": df["Biomarker adjustment"],
    "N": df["N"],
    "Matched-sample baseline β": df["base_beta_interaction"],
    "Biomarker-adjusted β": df["adjusted_beta_interaction"],
    "Adjusted 95% CI": [
        f"{lo:.4f} to {hi:.4f}"
        for lo, hi in zip(
            df["adjusted_CI_interaction_low"],
            df["adjusted_CI_interaction_high"],
        )
    ],
    "Adjusted P": df["adjusted_p_interaction"],
    "Adjusted q": df["q_adjusted_interaction"],
    "% change in |β|": df["attenuation_pct"],
})

CSV = OUT / "TableS25_BiomarkerAdjustment_PrimaryCardiacBrainInteractions.csv"
XLSX = OUT / "TableS25_BiomarkerAdjustment_PrimaryCardiacBrainInteractions.xlsx"

pub.to_csv(CSV, index=False)

# ------------------------------------------------------------
# Workbook helpers
# ------------------------------------------------------------

def style(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max(
            len(str(c.value)) if c.value is not None else 0
            for c in col
        )
        ws.column_dimensions[letter].width = min(width + 2, 55)

def write_df(ws, x):
    ws.append(list(x.columns))
    for row in x.itertuples(index=False, name=None):
        ws.append(list(row))
    style(ws)

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

# ------------------------------------------------------------
# Build workbook
# ------------------------------------------------------------

wb = Workbook()

ws = wb.active
ws.title = "Results"
write_df(ws, pub)

ws2 = wb.create_sheet("Numeric_Source")
write_df(ws2, num)

ws3 = wb.create_sheet("Summary")
write_df(ws3, summ)

prov_rows = []

for analysis, path in [
    ("publication source table", FULL),
    ("complete-case numeric model output", NUMERIC),
    ("four-association summary", SUMMARY),
    ("analysis script", ANALYSIS_SCRIPT),
    ("figure/table generation script", FIG_SCRIPT),
    ("Table S25 builder", BUILDER),
]:
    if path.exists():
        prov_rows.append({
            "analysis": analysis,
            "source": str(path),
            "sha256": sha(path),
        })

ws4 = wb.create_sheet("Provenance")
write_df(ws4, pd.DataFrame(prov_rows))

wb.save(XLSX)

# ------------------------------------------------------------
# Freeze sources/code
# ------------------------------------------------------------

for p in [FULL, NUMERIC, SUMMARY]:
    shutil.copy2(p, SRCDIR / p.name)

for p in [ANALYSIS_SCRIPT, FIG_SCRIPT, BUILDER]:
    if p.exists():
        shutil.copy2(p, CODEDIR / p.name)

# ------------------------------------------------------------
# README
# ------------------------------------------------------------

README = OUT / "README_PROVENANCE.txt"

README.write_text(
"""Supplementary Table S25. Plasma biomarker adjustment of the four primary cardiac–brain APOE ε4 interactions.

This table evaluates whether the four primary cardiac–brain interaction
estimates are attenuated after adjustment for circulating Alzheimer-related
and neurodegeneration biomarkers.

Four primary associations:
1. LVESVi × APOE ε4 → right superior frontal surface area
2. LV mass index × APOE ε4 → right lingual volume
3. LV mass index × APOE ε4 → right lingual surface area
4. LVEF × APOE ε4 → left middle temporal volume

Seven biomarker specifications:
1. Aβ42/40
2. p-tau181
3. GFAP
4. NfL
5. Aβ42/40 + p-tau181
6. GFAP + NfL
7. Aβ42/40 + p-tau181 + GFAP + NfL

Each adjusted model is compared with a baseline model fitted to the identical
complete-case sample. This prevents changes caused by biomarker missingness
from being mistaken for attenuation caused by biomarker adjustment.

All 28 adjusted interaction estimates remained directionally concordant with
their matched-sample baseline estimates and remained FDR-significant.

Observed percent changes in interaction magnitude were small.

This analysis is distinct from:
- Table S20: direct cardiac × APOE ε4 associations with plasma biomarkers
- Table S23: longitudinal cardiac trajectories predicting later biomarkers

Workbook sheets:
- Results
- Numeric_Source
- Summary
- Provenance
"""
)

# ------------------------------------------------------------
# Validation
# ------------------------------------------------------------

assert len(pub) == 28
assert len(pub["Primary cardiac–brain interaction"].unique()) == 4
assert len(pub["Biomarker adjustment"].unique()) == 7

# All adjusted q < .05
assert (df["q_adjusted_interaction"] < 0.05).all()

# Same direction baseline vs adjusted
same_direction = (
    (df["base_beta_interaction"] * df["adjusted_beta_interaction"]) > 0
)
assert same_direction.all()

# Summary file consistency
assert (summ["n_specs"] == 7).all()
assert (summ["n_same_direction"] == 7).all()
assert (summ["n_adjusted_q05"] == 7).all()

print("\nPUBLICATION TABLE:")
print(pub.to_string(index=False))

print("\nSUMMARY:")
print(summ.to_string(index=False))

print("\nMaximum absolute percent change in interaction magnitude:")
print(df["attenuation_pct"].abs().max())

print("\nAdjusted q range:")
print(df["q_adjusted_interaction"].min(),
      "to",
      df["q_adjusted_interaction"].max())

print("\nFILES:")
print(CSV)
print(XLSX)
print(README)

print(
    "\nSUCCESS: Table S25 created from the complete-case "
    "biomarker-adjustment robustness analysis."
)
