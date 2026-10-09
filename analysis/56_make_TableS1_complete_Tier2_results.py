#!/usr/bin/env python3

from pathlib import Path
import pandas as pd
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ============================================================
# PATHS
# ============================================================

INFILE = Path(
    "/data/qiallab/Framingham/results/"
    "tier2_corrected/tier2_corrected_all.csv"
)

OUTDIR = Path(
    "/data/qiallab/Framingham/results/"
    "paper_figures_tables/supplementary"
)
OUTDIR.mkdir(parents=True, exist_ok=True)

OUTXLSX = OUTDIR / "TableS1_complete_Tier2_results.xlsx"
OUTCSV  = OUTDIR / "TableS1_M3_APOE4_all_interactions.csv"

# ============================================================
# LOAD
# ============================================================

df = pd.read_csv(INFILE)

print(f"Loaded {len(df)} rows")
print("Columns:")
print(df.columns.tolist())

# ============================================================
# ADD PUBLICATION-FRIENDLY LABELS
# ============================================================

def clean_region(region):
    if pd.isna(region):
        return ""

    r = str(region)

    replacements = {
        "lh_": "Left ",
        "rh_": "Right ",
        "_vol": " volume",
        "_area": " surface area",
        "_tk": " thickness",
        "_tksd": " thickness variability",
        "_mcv": " MCV",
        "_": " ",
    }

    for old, new in replacements.items():
        r = r.replace(old, new)

    # readability cleanup
    r = r.replace("middletemporal", "middle temporal")
    r = r.replace("superiorfrontal", "superior frontal")
    r = r.replace("lingual", "lingual")

    return " ".join(r.split())


def clean_cardiac(x):
    mapping = {
        "LVEF": "LVEF",
        "LVEDVi": "LVEDVi",
        "LVESVi": "LVESVi",
        "LV_MASSi": "LV mass index",
    }
    return mapping.get(x, x)


df["Brain phenotype"] = df["region"].map(clean_region)
df["Cardiac phenotype"] = df["cardiac"].map(clean_cardiac)

# ============================================================
# ADD CI STRING
# ============================================================

def make_ci(lo, hi):
    if pd.isna(lo) or pd.isna(hi):
        return ""
    return f"{lo:.3f} to {hi:.3f}"


df["95% CI interaction"] = [
    make_ci(lo, hi)
    for lo, hi in zip(
        df["CI_interaction_low"],
        df["CI_interaction_high"]
    )
]

# ============================================================
# PUBLICATION COLUMN ORDER
# ============================================================

cols = [
    "metric_type",
    "family",
    "Brain phenotype",
    "Cardiac phenotype",
    "model",
    "N",
    "beta_cardiac",
    "SE_cardiac",
    "p_cardiac",
    "q_cardiac",
    "beta_interaction",
    "SE_interaction",
    "95% CI interaction",
    "p_interaction",
    "q_interaction",
    "R2",
]

# ============================================================
# CREATE SHEETS
# ============================================================

m3 = (
    df[df["model"] == "M3_APOE4"][cols]
    .copy()
    .sort_values(
        ["q_interaction", "p_interaction"],
        na_position="last"
    )
)

m2 = (
    df[df["model"] == "M2_vascular"][cols]
    .copy()
    .sort_values(
        ["q_cardiac", "p_cardiac"],
        na_position="last"
    )
)

m1 = (
    df[df["model"] == "M1_basic"][cols]
    .copy()
    .sort_values(
        ["q_cardiac", "p_cardiac"],
        na_position="last"
    )
)

sig = df[
    (
        (df["q_interaction"].notna()) &
        (df["q_interaction"] < 0.05)
    )
    |
    (
        (df["q_cardiac"].notna()) &
        (df["q_cardiac"] < 0.05)
    )
][cols].copy()

sig = sig.sort_values(
    ["model", "q_interaction", "q_cardiac"],
    na_position="last"
)

# Save M3 flat CSV too
m3.to_csv(OUTCSV, index=False)

# ============================================================
# NOTES / DATA DICTIONARY
# ============================================================

notes = pd.DataFrame({
    "Field": [
        "metric_type",
        "family",
        "Brain phenotype",
        "Cardiac phenotype",
        "model",
        "N",
        "beta_cardiac",
        "SE_cardiac",
        "p_cardiac",
        "q_cardiac",
        "beta_interaction",
        "SE_interaction",
        "95% CI interaction",
        "p_interaction",
        "q_interaction",
        "R2",
    ],
    "Definition": [
        "Brain morphometric outcome type.",
        "Anatomical family used for multiple-comparison correction.",
        "Publication-friendly brain region/outcome label.",
        "Publication-friendly cardiac phenotype label.",
        "Model specification: M1 basic, M2 vascular, or M3 APOE4 interaction.",
        "Complete-case sample size for the model.",
        "Standardized cardiac slope; in M3 this is the slope among APOE ε4 noncarriers.",
        "Robust standard error for the cardiac coefficient.",
        "Nominal P value for the cardiac coefficient.",
        "FDR-adjusted q value for the cardiac coefficient.",
        "Cardiac × APOE ε4 interaction coefficient.",
        "HC3 robust standard error for the interaction coefficient.",
        "95% confidence interval for the cardiac × APOE ε4 interaction.",
        "Nominal P value for the cardiac × APOE ε4 interaction.",
        "FDR-adjusted q value for the cardiac × APOE ε4 interaction.",
        "Model R².",
    ]
})

analysis_notes = pd.DataFrame({
    "Analysis note": [
        "Primary manuscript interaction model",
        "Reference group",
        "Standard errors",
        "Multiple-comparison correction",
        "Interpretation",
    ],
    "Details": [
        "M3_APOE4 models test cardiac phenotype × APOE ε4 carrier status.",
        "APOE ε4 noncarriers are the reference group.",
        "HC3 heteroskedasticity-robust standard errors were used.",
        "q values correspond to the prespecified FDR correction families used in the original Tier-2 analysis.",
        "Supplementary results should be interpreted in the context of the prespecified model hierarchy and manuscript primary hypotheses."
    ]
})

# ============================================================
# WRITE WORKBOOK
# ============================================================

with pd.ExcelWriter(OUTXLSX, engine="openpyxl") as writer:

    m3.to_excel(
        writer,
        sheet_name="M3_APOE4_interactions",
        index=False
    )

    m2.to_excel(
        writer,
        sheet_name="M2_vascular",
        index=False
    )

    m1.to_excel(
        writer,
        sheet_name="M1_basic",
        index=False
    )

    sig.to_excel(
        writer,
        sheet_name="FDR_significant",
        index=False
    )

    notes.to_excel(
        writer,
        sheet_name="Variable_dictionary",
        index=False,
        startrow=0
    )

    analysis_notes.to_excel(
        writer,
        sheet_name="Variable_dictionary",
        index=False,
        startrow=len(notes) + 3
    )

# ============================================================
# FORMAT WORKBOOK
# ============================================================

wb = load_workbook(OUTXLSX)

header_fill = PatternFill(
    "solid",
    fgColor="1F4E78"
)

header_font = Font(
    bold=True,
    color="FFFFFF"
)

sig_fill = PatternFill(
    "solid",
    fgColor="E2F0D9"
)

thin_gray = Side(
    style="thin",
    color="D9E2F3"
)

border = Border(
    bottom=thin_gray
)

for ws in wb.worksheets:

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Header
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    # General alignment
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )
            cell.border = border

    # Highlight FDR-significant rows
    header_map = {
        cell.value: idx + 1
        for idx, cell in enumerate(ws[1])
    }

    qint_col = header_map.get("q_interaction")
    qcard_col = header_map.get("q_cardiac")

    if qint_col or qcard_col:

        for r in range(2, ws.max_row + 1):

            significant = False

            if qint_col:
                val = ws.cell(r, qint_col).value
                if isinstance(val, (int, float)) and val < 0.05:
                    significant = True

            if qcard_col:
                val = ws.cell(r, qcard_col).value
                if isinstance(val, (int, float)) and val < 0.05:
                    significant = True

            if significant:
                for c in range(1, ws.max_column + 1):
                    ws.cell(r, c).fill = sig_fill

    # Column widths
    for col in range(1, ws.max_column + 1):

        letter = get_column_letter(col)
        header = ws.cell(1, col).value

        if header in ["Brain phenotype"]:
            width = 34

        elif header in [
            "Cardiac phenotype",
            "95% CI interaction"
        ]:
            width = 22

        elif header in [
            "metric_type",
            "family",
            "model"
        ]:
            width = 16

        elif header in ["Field"]:
            width = 24

        elif header in [
            "Definition",
            "Details",
            "Analysis note"
        ]:
            width = 70

        else:
            width = 14

        ws.column_dimensions[letter].width = width

    # Numeric formatting
    for row in ws.iter_rows(min_row=2):
        for cell in row:

            header = ws.cell(1, cell.column).value

            if header in [
                "beta_cardiac",
                "SE_cardiac",
                "beta_interaction",
                "SE_interaction",
                "R2",
            ]:
                cell.number_format = "0.000"

            elif header in [
                "p_cardiac",
                "q_cardiac",
                "p_interaction",
                "q_interaction",
            ]:
                cell.number_format = "0.000000"

# Save
wb.save(OUTXLSX)

# ============================================================
# SUMMARY
# ============================================================

print("\nSUPPLEMENTARY TABLE CREATED")
print("=" * 70)

print(f"M3 APOE4 interaction rows: {len(m3)}")
print(f"M2 vascular rows:         {len(m2)}")
print(f"M1 basic rows:            {len(m1)}")
print(f"FDR-significant rows:     {len(sig)}")

print("\nM3 APOE4 FDR-significant interactions:")
print(
    m3[m3["q_interaction"] < 0.05][
        [
            "Brain phenotype",
            "Cardiac phenotype",
            "N",
            "beta_interaction",
            "p_interaction",
            "q_interaction",
        ]
    ].to_string(index=False)
)

print("\nSaved:")
print(OUTXLSX)
print(OUTCSV)

