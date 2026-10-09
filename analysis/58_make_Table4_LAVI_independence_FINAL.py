#!/usr/bin/env python3

from pathlib import Path
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


# ============================================================
# PATHS
# ============================================================

INFILE = Path(
    "/data/qiallab/Framingham/results/"
    "atrial_APOE4/independence_LV/"
    "LAVI_APOE4_independence_LV_models.csv"
)

OUTDIR = Path(
    "/data/qiallab/Framingham/results/"
    "paper_figures_tables/atrial_APOE4"
)
OUTDIR.mkdir(parents=True, exist_ok=True)

OUTCSV = OUTDIR / "Table4_LAVI_APOE4_independence_FINAL.csv"
OUTXLSX = OUTDIR / "Table4_LAVI_APOE4_independence_FINAL.xlsx"
OUTPNG = OUTDIR / "Table4_LAVI_APOE4_independence_FINAL.png"

# Supplementary all-sample sensitivity
OUTSCSV = OUTDIR / "TableS2_LAVI_APOE4_independence_all_sample.csv"


# ============================================================
# LOAD EXACT SOURCE
# ============================================================

df = pd.read_csv(INFILE)

print("\nSOURCE COLUMNS:")
print(df.columns.tolist())

print("\nSOURCE DATA:")
print(df.to_string(index=False))


# ============================================================
# REQUIRE EXACT SOURCE SCHEMA
# ============================================================

required = [
    "analysis",
    "model",
    "outcome",
    "atrial_metric",
    "N",
    "N_APOE4_noncarrier",
    "N_APOE4_carrier",
    "beta_noncarrier",
    "SE_noncarrier",
    "beta_interaction",
    "SE_interaction",
    "CI_interaction_low",
    "CI_interaction_high",
    "p_interaction",
    "beta_carrier",
    "SE_carrier",
    "CI_carrier_low",
    "CI_carrier_high",
    "beta_LV_MASSi",
    "p_LV_MASSi",
    "beta_LVEF",
    "p_LVEF",
    "R2",
    "interaction_abs_change_pct",
]

missing = [c for c in required if c not in df.columns]

if missing:
    raise RuntimeError(
        f"Missing required columns: {missing}"
    )


# ============================================================
# VALIDATE EXPECTED ANALYSIS
# ============================================================

if set(df["outcome"].dropna().unique()) != {
    "lh_G_occipital_middle_tksd"
}:
    raise RuntimeError(
        "Unexpected atrial outcome in source file."
    )

if set(df["atrial_metric"].dropna().unique()) != {
    "LAVI_max"
}:
    raise RuntimeError(
        "Unexpected atrial metric in source file."
    )


# ============================================================
# MODEL LABELS
# ============================================================

model_labels = {
    "M0_ATRIAL_BASE":
        "Baseline",

    "M1_PLUS_LV_MASSi":
        "+ LV mass index",

    "M2_PLUS_LVEF":
        "+ LVEF",

    "M3_PLUS_LV_MASSi_LVEF":
        "+ LV mass index + LVEF",
}

model_order = [
    "M0_ATRIAL_BASE",
    "M1_PLUS_LV_MASSi",
    "M2_PLUS_LVEF",
    "M3_PLUS_LV_MASSi_LVEF",
]


# ============================================================
# FUNCTION TO BUILD PUBLICATION TABLE
# ============================================================

def make_publication_table(d):

    d = d.copy()

    d["model_order"] = pd.Categorical(
        d["model"],
        categories=model_order,
        ordered=True
    )

    d = (
        d.sort_values("model_order")
        .drop(columns="model_order")
        .reset_index(drop=True)
    )

    if len(d) != 4:
        raise RuntimeError(
            f"Expected 4 nested models, found {len(d)}"
        )

    pub = pd.DataFrame({
        "Model":
            d["model"].map(model_labels),

        "N":
            d["N"].astype(int),

        "APOE ε4 noncarriers":
            d["N_APOE4_noncarrier"].astype(int),

        "APOE ε4 carriers":
            d["N_APOE4_carrier"].astype(int),

        "β noncarrier":
            d["beta_noncarrier"],

        "β interaction":
            d["beta_interaction"],

        "95% CI":
            [
                f"{lo:.3f} to {hi:.3f}"
                for lo, hi in zip(
                    d["CI_interaction_low"],
                    d["CI_interaction_high"]
                )
            ],

        "P interaction":
            d["p_interaction"],

        "β carrier":
            d["beta_carrier"],

        "R²":
            d["R2"],
    })

    return pub


# ============================================================
# PRIMARY EXAM-8 TABLE 4
# ============================================================

primary = df[
    df["analysis"] == "PRIMARY_EXAM8"
].copy()

pub = make_publication_table(primary)

print("\nFINAL TABLE 4 — PRIMARY EXAM 8")
print("=" * 125)
print(pub.to_string(index=False))


# ============================================================
# SENSITIVITY ALL-SAMPLE TABLE
# ============================================================

sens = df[
    df["analysis"] == "SENSITIVITY_ALL"
].copy()

pub_sens = make_publication_table(sens)

print("\nSUPPLEMENTARY ALL-SAMPLE SENSITIVITY")
print("=" * 125)
print(pub_sens.to_string(index=False))


# ============================================================
# NUMERICAL VALIDATION CHECKS
# ============================================================

expected_primary_beta = [
    -0.544973,
    -0.549027,
    -0.544608,
    -0.548698,
]

observed = primary.set_index("model").loc[
    model_order,
    "beta_interaction"
].values

if not np.allclose(
    observed,
    expected_primary_beta,
    atol=1e-6
):
    raise RuntimeError(
        "Primary interaction coefficients differ from "
        "the validated values."
    )

if not (
    primary["N"].eq(451).all()
    and primary["N_APOE4_noncarrier"].eq(355).all()
    and primary["N_APOE4_carrier"].eq(96).all()
):
    raise RuntimeError(
        "Primary Exam-8 sample counts differ from expected."
    )

print("\nVALIDATION CHECKS PASSED:")
print("  Primary N = 451")
print("  APOE ε4 noncarriers = 355")
print("  APOE ε4 carriers = 96")
print("  All four validated β interaction values reproduced.")


# ============================================================
# SAVE CSV
# ============================================================

pub.to_csv(
    OUTCSV,
    index=False
)

pub_sens.to_csv(
    OUTSCSV,
    index=False
)


# ============================================================
# SAVE XLSX
# ============================================================

with pd.ExcelWriter(
    OUTXLSX,
    engine="openpyxl"
) as writer:

    pub.to_excel(
        writer,
        sheet_name="Table 4",
        index=False,
        startrow=2
    )

    pub_sens.to_excel(
        writer,
        sheet_name="Sensitivity all sample",
        index=False,
        startrow=2
    )


# ============================================================
# FORMAT XLSX
# ============================================================

wb = load_workbook(OUTXLSX)

header_fill = PatternFill(
    "solid",
    fgColor="1F4E78"
)

header_font = Font(
    bold=True,
    color="FFFFFF"
)

thin = Side(
    style="thin",
    color="D9E2F3"
)


# ------------------------------------------------------------
# MAIN TABLE
# ------------------------------------------------------------

ws = wb["Table 4"]

title = (
    "Table 4. LAVI × APOE ε4 interaction after adjustment "
    "for ventricular structure and function"
)

ws["A1"] = title
ws.merge_cells("A1:J1")

ws["A1"].font = Font(
    bold=True,
    size=14
)

ws["A1"].alignment = Alignment(
    horizontal="center",
    vertical="center"
)

for cell in ws[3]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )

for row in ws.iter_rows(
    min_row=4,
    max_row=7
):
    for cell in row:
        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True
        )
        cell.border = Border(
            bottom=thin
        )

# Left-align model labels
for row in range(4, 8):
    ws.cell(row, 1).alignment = Alignment(
        horizontal="left",
        vertical="center",
        wrap_text=True
    )

# Column widths
widths = {
    "A": 30,
    "B": 9,
    "C": 18,
    "D": 16,
    "E": 15,
    "F": 15,
    "G": 22,
    "H": 16,
    "I": 14,
    "J": 10,
}

for col, width in widths.items():
    ws.column_dimensions[col].width = width

# Numeric formatting
for row in range(4, 8):

    # β noncarrier
    ws.cell(row, 5).number_format = "0.000"

    # β interaction
    ws.cell(row, 6).number_format = "0.000"

    # p interaction
    ws.cell(row, 8).number_format = "0.00E+00"

    # β carrier
    ws.cell(row, 9).number_format = "0.000"

    # R²
    ws.cell(row, 10).number_format = "0.000"


# ------------------------------------------------------------
# MAIN TABLE FOOTNOTES
# ------------------------------------------------------------

notes = [
    (
        "Outcome: left middle-occipital cortical thickness "
        "variability. Exposure: maximum left-atrial volume "
        "index (LAVI)."
    ),
    (
        "β noncarrier is the standardized LAVI–brain slope "
        "among APOE ε4 noncarriers. β interaction represents "
        "the difference in slope among APOE ε4 carriers."
    ),
    (
        "β carrier = β noncarrier + β interaction."
    ),
    (
        "The four nested models evaluate whether the "
        "LAVI × APOE ε4 interaction persists after adjustment "
        "for LV mass index, LVEF, or both."
    ),
    (
        "Interaction P values are nominal targeted robustness "
        "tests; the original LAVI × APOE ε4 discovery result "
        "survived FDR correction in the atrial screening analysis."
    ),
]

for i, note in enumerate(
    notes,
    start=9
):

    ws[f"A{i}"] = note

    ws.merge_cells(
        start_row=i,
        start_column=1,
        end_row=i,
        end_column=10
    )

    ws[f"A{i}"].font = Font(
        italic=True,
        size=9
    )

    ws[f"A{i}"].alignment = Alignment(
        wrap_text=True
    )

ws.freeze_panes = "A4"


# ------------------------------------------------------------
# SENSITIVITY SHEET
# ------------------------------------------------------------

ws2 = wb["Sensitivity all sample"]

ws2["A1"] = (
    "Supplementary sensitivity: LAVI × APOE ε4 "
    "nested ventricular-adjustment models in all available "
    "participants"
)

ws2.merge_cells("A1:J1")

ws2["A1"].font = Font(
    bold=True,
    size=14
)

ws2["A1"].alignment = Alignment(
    horizontal="center"
)

for cell in ws2[3]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )

for row in ws2.iter_rows(
    min_row=4,
    max_row=7
):
    for cell in row:
        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True
        )
        cell.border = Border(
            bottom=thin
        )

for row in range(4, 8):
    ws2.cell(row, 1).alignment = Alignment(
        horizontal="left",
        vertical="center",
        wrap_text=True
    )

for col, width in widths.items():
    ws2.column_dimensions[col].width = width

for row in range(4, 8):
    ws2.cell(row, 5).number_format = "0.000"
    ws2.cell(row, 6).number_format = "0.000"
    ws2.cell(row, 8).number_format = "0.00E+00"
    ws2.cell(row, 9).number_format = "0.000"
    ws2.cell(row, 10).number_format = "0.000"

ws2.freeze_panes = "A4"

wb.save(OUTXLSX)


# ============================================================
# PUBLICATION PNG — MAIN TABLE
# ============================================================

display = pub.copy()

for c in [
    "β noncarrier",
    "β interaction",
    "β carrier",
    "R²",
]:
    display[c] = display[c].map(
        lambda x:
        f"{x:.3f}"
        if pd.notna(x)
        else ""
    )

display["P interaction"] = display[
    "P interaction"
].map(
    lambda x:
    f"{x:.2e}"
    if pd.notna(x)
    else ""
)

fig, ax = plt.subplots(
    figsize=(15.5, 4.8)
)

ax.axis("off")

tbl = ax.table(
    cellText=display.values,
    colLabels=display.columns,
    cellLoc="center",
    colLoc="center",
    loc="center"
)

tbl.auto_set_font_size(False)
tbl.set_fontsize(9.2)
tbl.scale(1, 1.7)

for c in range(
    len(display.columns)
):
    tbl[(0, c)].set_text_props(
        weight="bold"
    )

for r in range(
    1,
    len(display) + 1
):
    tbl[(r, 0)].get_text().set_ha(
        "left"
    )

ax.set_title(
    title,
    fontsize=13,
    fontweight="bold",
    pad=14
)

fig.text(
    0.01,
    0.015,
    "Outcome: left middle-occipital cortical thickness "
    "variability. LAVI × APOE ε4 interaction estimates remain "
    "essentially unchanged after adjustment for LV mass index, "
    "LVEF, or both.",
    fontsize=8.5,
    wrap=True
)

plt.tight_layout(
    rect=[0, 0.06, 1, 0.94]
)

plt.savefig(
    OUTPNG,
    dpi=300,
    bbox_inches="tight"
)

plt.close()


# ============================================================
# FINAL NO-NaN CHECK ON MANUSCRIPT TABLE
# ============================================================

check_cols = [
    "Model",
    "N",
    "APOE ε4 noncarriers",
    "APOE ε4 carriers",
    "β noncarrier",
    "β interaction",
    "95% CI",
    "P interaction",
    "β carrier",
    "R²",
]

if pub[check_cols].isna().any().any():

    print("\nWARNING — NaNs detected:")
    print(
        pub[check_cols]
        .isna()
        .sum()
    )

else:
    print(
        "\nFINAL CHECK: no NaNs in Table 4."
    )


# ============================================================
# SAVE SUMMARY
# ============================================================

print("\nSaved:")
print(OUTCSV)
print(OUTXLSX)
print(OUTPNG)
print(OUTSCSV)

