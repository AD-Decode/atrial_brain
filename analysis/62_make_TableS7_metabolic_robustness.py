#!/usr/bin/env python3

from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path("/data/qiallab/Framingham")

INDIR = ROOT / "results" / "primary_pairs_metabolic_modifiers"

OUTDIR = (
    ROOT / "results" /
    "paper_figures_tables" /
    "SUBMISSION_PACKAGE" /
    "04_Supplementary_Tables"
)

OUTDIR.mkdir(parents=True, exist_ok=True)

OUTFILE = OUTDIR / "TableS7_metabolic_robustness.xlsx"

ATTEN_FILE = INDIR / "primary_pairs_metabolic_attenuation.csv"
JOINT_FILE = INDIR / "primary_pairs_metabolic_joint_attenuation.csv"
THREE_FILE = INDIR / "primary_pairs_metabolic_threeway.csv"

for f in [ATTEN_FILE, JOINT_FILE, THREE_FILE]:
    if not f.exists():
        raise FileNotFoundError(f"Missing required input: {f}")

atten = pd.read_csv(ATTEN_FILE)
joint = pd.read_csv(JOINT_FILE)
three = pd.read_csv(THREE_FILE)

# ------------------------------------------------------------
# Helpers
# ------------------------------------------------------------

def fmt_ci(lo, hi):
    return f"{lo:.3f} to {hi:.3f}"

def fmt_p(p):
    if pd.isna(p):
        return ""
    if p < 0.0001:
        return f"{p:.2e}"
    return f"{p:.4f}"

def fmt_beta(x):
    if pd.isna(x):
        return ""
    return f"{x:.3f}"

def fmt_pct(x):
    if pd.isna(x):
        return ""
    return f"{x:.1f}"

# ------------------------------------------------------------
# Prepare Panel A
# ------------------------------------------------------------

panel_a = pd.DataFrame({
    "Association": atten["association"],
    "Modifier": atten["modifier"],
    "N": atten["N_matched"].astype(int),
    "APOE ε4 noncarriers, n": atten["N_APOE4_noncarrier"].astype(int),
    "APOE ε4 carriers, n": atten["N_APOE4_carrier"].astype(int),
    "Matched M3 interaction β": atten["beta_interaction_M3_matched"].map(fmt_beta),
    "M3 + metabolic marker interaction β": atten["beta_interaction_plus_metabolic"].map(fmt_beta),
    "Change in |β| vs matched M3, %": atten["attenuation_abs_beta_percent"].map(fmt_pct),
    "P after metabolic adjustment": atten["p_plus_metabolic"].map(fmt_p),
})

# ------------------------------------------------------------
# Prepare Panel B
# ------------------------------------------------------------

panel_b = pd.DataFrame({
    "Association": joint["association"],
    "N": joint["N_common"].astype(int),
    "APOE ε4 noncarriers, n": joint["N_APOE4_noncarrier"].astype(int),
    "APOE ε4 carriers, n": joint["N_APOE4_carrier"].astype(int),
    "Matched M3 interaction β": joint["beta_interaction_M3_matched"].map(fmt_beta),
    "Matched M3 95% CI": [
        fmt_ci(lo, hi)
        for lo, hi in zip(
            joint["CI_low_M3_matched"],
            joint["CI_high_M3_matched"]
        )
    ],
    "Matched M3 P": joint["p_M3_matched"].map(fmt_p),
    "M3 + HbA1c + insulin + adiponectin interaction β":
        joint["beta_interaction_plus_all3"].map(fmt_beta),
    "Adjusted 95% CI": [
        fmt_ci(lo, hi)
        for lo, hi in zip(
            joint["CI_low_plus_all3"],
            joint["CI_high_plus_all3"]
        )
    ],
    "Adjusted P": joint["p_plus_all3"].map(fmt_p),
    "Change in |β| vs matched M3, %":
        joint["attenuation_abs_beta_percent"].map(fmt_pct),
})

# ------------------------------------------------------------
# Prepare three-way tests
# ------------------------------------------------------------

three = three.sort_values(
    ["p_threeway", "association", "modifier"]
).reset_index(drop=True)

three_sheet = pd.DataFrame({
    "Association": three["association"],
    "Modifier": three["modifier"],
    "N": three["N"].astype(int),
    "APOE ε4 noncarriers, n": three["N_APOE4_noncarrier"].astype(int),
    "APOE ε4 carriers, n": three["N_APOE4_carrier"].astype(int),
    "Three-way interaction β": three["beta_threeway"].map(fmt_beta),
    "95% CI": [
        fmt_ci(lo, hi)
        for lo, hi in zip(
            three["CI_low"],
            three["CI_high"]
        )
    ],
    "P": three["p_threeway"].map(fmt_p),
    "FDR q": three["q_threeway_FDR12"].map(fmt_p),
})

# ------------------------------------------------------------
# Workbook
# ------------------------------------------------------------

wb = Workbook()

ws = wb.active
ws.title = "Table S7"

ws2 = wb.create_sheet(
    "Three-way tests"
)

# Styles
title_font = Font(
    bold=True,
    size=14
)

section_font = Font(
    bold=True,
    size=11
)

header_font = Font(
    bold=True,
    size=10
)

header_fill = PatternFill(
    "solid",
    fgColor="D9EAF7"
)

thin_gray = Side(
    style="thin",
    color="B7B7B7"
)

header_border = Border(
    bottom=thin_gray
)

wrap_top = Alignment(
    wrap_text=True,
    vertical="top"
)

center_wrap = Alignment(
    horizontal="center",
    vertical="center",
    wrap_text=True
)

# ------------------------------------------------------------
# Main sheet title
# ------------------------------------------------------------

ws["A1"] = (
    "Table S7. Metabolic robustness of primary "
    "cardiac × APOE ε4 heart–brain associations"
)
ws["A1"].font = title_font

ws.merge_cells(
    start_row=1,
    start_column=1,
    end_row=1,
    end_column=11
)

# ------------------------------------------------------------
# Panel A
# ------------------------------------------------------------

ws["A3"] = (
    "Panel A. Matched-sample attenuation after adjustment "
    "for individual metabolic markers"
)
ws["A3"].font = section_font

row0 = 4

for col_idx, col in enumerate(
    panel_a.columns,
    start=1
):
    cell = ws.cell(
        row=row0,
        column=col_idx,
        value=col
    )
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center_wrap
    cell.border = header_border

for i, row in panel_a.iterrows():
    excel_row = row0 + 1 + i

    for j, value in enumerate(
        row.tolist(),
        start=1
    ):
        cell = ws.cell(
            row=excel_row,
            column=j,
            value=value
        )

        cell.alignment = (
            wrap_top
            if j <= 2
            else center_wrap
        )

# ------------------------------------------------------------
# Panel B
# ------------------------------------------------------------

row_b_title = row0 + len(panel_a) + 3

ws.cell(
    row=row_b_title,
    column=1,
    value=(
        "Panel B. Joint adjustment for HbA1c, insulin, "
        "and adiponectin on a common complete-case sample"
    )
).font = section_font

row_b_header = row_b_title + 1

for col_idx, col in enumerate(
    panel_b.columns,
    start=1
):
    cell = ws.cell(
        row=row_b_header,
        column=col_idx,
        value=col
    )
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center_wrap
    cell.border = header_border

for i, row in panel_b.iterrows():
    excel_row = row_b_header + 1 + i

    for j, value in enumerate(
        row.tolist(),
        start=1
    ):
        cell = ws.cell(
            row=excel_row,
            column=j,
            value=value
        )

        cell.alignment = (
            wrap_top
            if j == 1
            else center_wrap
        )

# ------------------------------------------------------------
# Notes
# ------------------------------------------------------------

notes_row = (
    row_b_header
    + len(panel_b)
    + 2
)

note = (
    "Notes: Outcomes and cardiac predictors were standardized. "
    "Models used HC3 robust standard errors and the established "
    "M3 covariates: age, sex, MRI–CMR interval, BMI, systolic "
    "blood pressure, current smoking, diabetes history, and total "
    "cortical surface area for surface-area outcomes. Insulin and "
    "adiponectin were log-transformed before standardization. "
    "Change in |β| is calculated relative to the matched-sample "
    "M3 interaction estimate; negative values indicate a slightly "
    "larger absolute interaction after metabolic adjustment. "
    "These analyses assess robustness/attenuation and should not "
    "be interpreted as mediation."
)

ws.cell(
    row=notes_row,
    column=1,
    value=note
)

ws.merge_cells(
    start_row=notes_row,
    start_column=1,
    end_row=notes_row + 2,
    end_column=11
)

ws.cell(
    row=notes_row,
    column=1
).alignment = wrap_top

# ------------------------------------------------------------
# Secondary sheet
# ------------------------------------------------------------

ws2["A1"] = (
    "Prespecified cardiac × APOE ε4 × metabolic-marker "
    "interaction tests"
)
ws2["A1"].font = title_font

ws2.merge_cells(
    start_row=1,
    start_column=1,
    end_row=1,
    end_column=9
)

header_row = 3

for col_idx, col in enumerate(
    three_sheet.columns,
    start=1
):
    cell = ws2.cell(
        row=header_row,
        column=col_idx,
        value=col
    )
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center_wrap
    cell.border = header_border

for i, row in three_sheet.iterrows():
    excel_row = header_row + 1 + i

    for j, value in enumerate(
        row.tolist(),
        start=1
    ):
        cell = ws2.cell(
            row=excel_row,
            column=j,
            value=value
        )

        cell.alignment = (
            wrap_top
            if j <= 2
            else center_wrap
        )

note_row = (
    header_row
    + len(three_sheet)
    + 2
)

ws2.cell(
    row=note_row,
    column=1,
    value=(
        "FDR correction was applied jointly across 12 prespecified "
        "three-way tests (4 primary heart–brain pairs × 3 metabolic "
        "modifiers). No interaction survived FDR correction."
    )
)

ws2.merge_cells(
    start_row=note_row,
    start_column=1,
    end_row=note_row + 1,
    end_column=9
)

ws2.cell(
    row=note_row,
    column=1
).alignment = wrap_top

# ------------------------------------------------------------
# Column widths
# ------------------------------------------------------------

main_widths = {
    "A": 48,
    "B": 16,
    "C": 11,
    "D": 18,
    "E": 16,
    "F": 19,
    "G": 24,
    "H": 20,
    "I": 18,
    "J": 18,
    "K": 20,
}

for col, width in main_widths.items():
    ws.column_dimensions[col].width = width

three_widths = {
    "A": 48,
    "B": 16,
    "C": 10,
    "D": 18,
    "E": 16,
    "F": 20,
    "G": 20,
    "H": 13,
    "I": 13,
}

for col, width in three_widths.items():
    ws2.column_dimensions[col].width = width

# ------------------------------------------------------------
# Freeze panes
# ------------------------------------------------------------

ws.freeze_panes = "A5"
ws2.freeze_panes = "A4"

# ------------------------------------------------------------
# Row heights
# ------------------------------------------------------------

ws.row_dimensions[1].height = 24
ws.row_dimensions[4].height = 42
ws.row_dimensions[row_b_header].height = 50

ws2.row_dimensions[1].height = 24
ws2.row_dimensions[3].height = 42

# ------------------------------------------------------------
# Save
# ------------------------------------------------------------

wb.save(OUTFILE)

print("\nSaved Table S7:")
print(OUTFILE)

print("\nPanel A rows:", len(panel_a))
print("Panel B rows:", len(panel_b))
print("Three-way rows:", len(three_sheet))

print("\nJoint-adjustment results:")
print(
    panel_b[
        [
            "Association",
            "N",
            "M3 + HbA1c + insulin + adiponectin interaction β",
            "Adjusted P",
            "Change in |β| vs matched M3, %",
        ]
    ].to_string(index=False)
)

print("\nDone.")
